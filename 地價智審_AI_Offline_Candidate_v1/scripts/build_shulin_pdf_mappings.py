# -*- coding: utf-8 -*-
"""
build_shulin_pdf_mappings.py — OFFICIAL-SIX-PAGE-PDF-E1 Task 3.

Build-time (PyMuPDF-only, no pywin32/Excel dependency -- safe to re-run
anywhere the repo is checked out) generator for the 3 auditable overlay
mapping configs:

  data/templates/shulin/shulin_table3_mapping.json
  data/templates/shulin/shulin_table51_mapping.json
  data/templates/shulin/shulin_table4_mapping.json

Every entry traces back to a REAL cell/merged-range reference in the
corresponding official XLSX sheet (never a hand-guessed pixel number):
this script reads each field's `source_cell` (an Excel A1-style range,
taken directly from the openpyxl audit of the 3 official workbooks --
see docs/phase7/official_six_page_pdf_e1.md Task 1/3) and resolves it to
a PDF-point rectangle via pdf/shulin_pdf_grid.py's grid-boundary
extraction against the ALREADY-BUILT blank template PDF (data/templates/
shulin/shulin_table{3,51,4}_blank_v1.pdf, produced by
build_shulin_official_pdf_templates.py) -- i.e. the rectangle is
measured from the real rendered grid Excel itself drew, not computed
from an Excel-column-width-to-pixel formula.

Fields with NO corresponding data source in this competition's runtime
(e.g. Table3's facility-name/distance checkboxes -- schools, markets,
funeral facilities, major stations -- for which Shulin has no confirmed-
facility-selection data source, unlike the legacy Jinshan flow's
facility_confirmation_repository.py) are simply absent from the catalogs
below. This is deliberate (never guessed, never fabricated) and is
called out explicitly in the E1 report as a known, honest scope limit.
"""
from __future__ import annotations

import json
import os

import fitz
import openpyxl
from openpyxl.utils import range_boundaries, get_column_letter

import sys
_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.dirname(_THIS_DIR)
sys.path.insert(0, _REPO_ROOT)
sys.path.insert(0, os.path.join(_REPO_ROOT, "pdf"))
from shulin_pdf_grid import extract_grid_boundaries, rect_for_span, reconcile_boundaries  # noqa: E402

_TEMPLATE_DIR = os.path.join(_REPO_ROOT, "data", "templates", "shulin")
_SRC_DIR = os.path.join(_REPO_ROOT, "data", "sources", "competition", "shulin_residential_2026")

DEFAULT_FONT_SIZE = 8.5
DEFAULT_MIN_FONT_SIZE = 5.5

_DEFAULT_COL_WIDTH_CHARS = 8.43  # Excel's own fallback when a column has no explicit width


def _sheet_grid(xlsx_name: str, sheet_name: str, blank_pdf_name: str, num_cols: int, num_rows: int,
                 force_leading_rows: "int | None" = None):
    """Returns (x_boundaries, y_boundaries), reconciled to EXACTLY
    num_cols+1 / num_rows+1 entries using the sheet's own declared column
    widths / row heights as the proportional reference (see
    reconcile_boundaries()'s docstring for why this is necessary -- 表3 in
    particular has a few genuinely borderless spacer columns that
    get_drawings()-based detection alone cannot find). `force_leading_rows`
    is passed straight to reconcile_boundaries()'s `force_leading` for the
    row axis only -- see that function's docstring for why 表3 needs this
    (rows 1-2 have no cell borders at all, and the auto-fit alone was
    verified, via page.search_for() ground truth, to sometimes pick the
    wrong neighboring split)."""
    ws = openpyxl.load_workbook(os.path.join(_SRC_DIR, xlsx_name))[sheet_name]
    col_widths = []
    for i in range(1, num_cols + 1):
        letter = get_column_letter(i)
        dim = ws.column_dimensions.get(letter)
        col_widths.append(dim.width if (dim and dim.width) else _DEFAULT_COL_WIDTH_CHARS)
    default_row_h = ws.sheet_format.defaultRowHeight or 15.0
    row_heights = []
    for r in range(1, num_rows + 1):
        dim = ws.row_dimensions.get(r)
        row_heights.append(dim.height if (dim and dim.height) else default_row_h)

    doc = fitz.open(os.path.join(_TEMPLATE_DIR, blank_pdf_name))
    raw_xs, raw_ys = extract_grid_boundaries(doc[0])
    doc.close()
    xs = reconcile_boundaries(raw_xs, col_widths)
    ys = reconcile_boundaries(raw_ys, row_heights, force_leading=force_leading_rows)
    return xs, ys


def _resolve(xs, ys, source_cell: str):
    min_col, min_row, max_col, max_row = range_boundaries(source_cell)
    return rect_for_span(xs, ys, min_col, min_row, max_col, max_row)


def _entry(field_id, source_sheet, source_cell, xs, ys, *, align="left",
           font_size=DEFAULT_FONT_SIZE, min_font_size=DEFAULT_MIN_FONT_SIZE,
           formatter="str", kind="text", left_pad=None):
    entry = {
        "field_id": field_id,
        "source_sheet": source_sheet,
        "source_cell": source_cell,
        "rect": _resolve(xs, ys, source_cell),
        "font_size": font_size,
        "min_font_size": min_font_size,
        "align": align,
        "formatter": formatter,
        "kind": kind,
    }
    if left_pad is not None:
        entry["left_pad"] = left_pad
    return entry


# ---------------------------------------------------------------------------
# TABLE3 (表3區段勘查表) -- one field catalog, reused for all 4 pages
# (P002/P003/P004/P001), since every page overlays the SAME blank template
# with that page's own segment's data.
# ---------------------------------------------------------------------------

def build_table3_mapping() -> list:
    sheet = "表3區段勘查表"
    xs, ys = _sheet_grid(
        "表3地價區段勘查表.xlsx", sheet, "shulin_table3_blank_v1.pdf", num_cols=22, num_rows=46,
        # Ground-truth verified (page.search_for() against the real blank
        # template): rows 1-2 (title + column-header row) have NO cell
        # borders anywhere, so the first REAL detected row boundary is
        # already row 3's own top edge -- i.e. exactly 2 leading row
        # boundaries are missing, not whatever the auto-fit alone guesses.
        force_leading_rows=2,
    )

    entries = [
        _entry("year_period", sheet, "B3:D3", xs, ys, align="center"),
        _entry("segment_code", sheet, "G3:H3", xs, ys, align="center"),
        _entry("zone_range_description", sheet, "L3:V3", xs, ys, font_size=7, min_font_size=5),
        # NOTE (fixed after visual overlay verification): D{row}:G4-style
        # cells hold the row's own LABEL text (e.g. D4:G4="都市計畫(內外)"),
        # NOT the fillable value -- the real value cell is the NEXT column
        # group over (H4:K4 for rows 4-9/10, I{row}:K{row} for rows 23-29,
        # confirmed cell-by-cell against the real XLSX). H6:K6/H7:K7 also
        # carry a static "-" placeholder glyph already, so those two use
        # kind="redact_and_write" (redact the existing glyph before writing
        # the real value) rather than a plain overlay that would collide
        # with it.
        # left_pad is widened for rows whose own LABEL text visually runs
        # close to (or slightly past) its nominal cell edge -- confirmed via
        # rendered-page visual inspection, not guessed; the value would
        # otherwise sit flush against the tail of the label with no visible
        # gap between them.
        # left_pad values below are measured directly (page.search_for() on
        # the real blank template) as this row's OWN label's rendered
        # right edge minus the value rect's own x0, +~4pt clearance -- e.g.
        # "都市計畫(內外)" label span ends at x=97.7, H4:K4's x0=84.6, so a
        # value written flush at x0 would visually collide with the tail
        # of the label; not a wrong-cell placement, just insufficient
        # clearance from a label wider than its own nominal D:G span.
        _entry("regional_zoning_inside_outside", sheet, "H4:K4", xs, ys, left_pad=17),
        _entry("regional_land_use_zone", sheet, "H5:K5", xs, ys, left_pad=34),
        _entry("regional_building_coverage_ratio", sheet, "H6:K6", xs, ys, align="right", formatter="pct_no_unit", kind="redact_and_write"),
        _entry("regional_floor_area_ratio", sheet, "H7:K7", xs, ys, align="right", formatter="pct_no_unit", kind="redact_and_write"),
        _entry("regional_construction_prohibited", sheet, "H8:K8", xs, ys, left_pad=12),
        _entry("regional_construction_restricted", sheet, "H9:K10", xs, ys, font_size=7.5),
        _entry("regional_road_development_level", sheet, "I23:K23", xs, ys, left_pad=32, font_size=6.5, min_font_size=4.5),
        _entry("regional_sunlight", sheet, "I24:K24", xs, ys),
        _entry("regional_view", sheet, "I25:K25", xs, ys),
        _entry("regional_slope", sheet, "I26:K26", xs, ys),
        _entry("regional_drainage_quality", sheet, "I27:K27", xs, ys, left_pad=41),
        _entry("regional_terrain", sheet, "I28:K28", xs, ys),
        # NOTE (visual overlay verification, docs/phase7/official_six_page_
        # pdf_e1.md Task 3): main_road_name/main_road_width (row 11) and
        # building_density_pct/building_type (rows 42-43) sit in a densely
        # merged, irregular region of this sheet (rows 35-44 mix multi-row
        # spans like A35:A44/L40:N41/Q40:V41) where the row-height-
        # proportional reconciliation in reconcile_boundaries() could not be
        # visually confirmed as pixel-accurate -- rather than risk
        # overlaying real text in the WRONG cell, these 4 fields are
        # deliberately left OUT of this catalog (blank on the rendered PDF)
        # until a follow-up round can re-verify their exact boundaries.
        # land_use_status is a multi-select checkbox line, redrawn (never a
        # plain text overlay) by shulin_official_pdf_renderer.py -- kind
        # "checkbox_multiselect" marks it for that special handling; the
        # rect below is the FULL line's own bbox for redaction+redraw.
        _entry("land_use_status", sheet, "Q44:V44", xs, ys, font_size=8, kind="checkbox_multiselect"),
    ]
    return entries


# ---------------------------------------------------------------------------
# TABLE51 (表5-1區域因素明細表(住)) -- 29 regional factors across 8
# categories + 8 category subtotals + 1 grand total, each with base grade
# (no pct -- base is the reference) + 3 comparables' grade+pct.
# ---------------------------------------------------------------------------

# (category_index, factor field_id, factor row) -- one row per factor, in
# official-form order; matches data/rules/competition/shulin_residential_2026/
# regional_rules.json's own 29 REG-*/regional_other catalog (SHULIN-
# COMPETITION-RULE-PACK-A2), cross-verified row-by-row against the real
# 表5-1區域因素明細表(住) sheet (docs/phase7/official_six_page_pdf_e1.md Task 1).
_TABLE51_FACTOR_ROWS = [
    (1, "regional_zoning_inside_outside", 5),
    (1, "regional_land_use_zone", 6),
    (1, "regional_building_coverage_ratio", 7),
    (1, "regional_floor_area_ratio", 8),
    (1, "regional_construction_prohibited", 9),
    (1, "regional_construction_restricted", 10),
    (2, "regional_main_road_width", 12),
    (2, "regional_avg_road_width", 13),
    (2, "regional_major_station_proximity", 14),
    (2, "regional_bus_stop_proximity", 15),
    (2, "regional_interchange_proximity", 16),
    (2, "regional_road_development_level", 17),
    (3, "regional_sunlight", 19),
    (3, "regional_view", 20),
    (3, "regional_slope", 21),
    (3, "regional_drainage_quality", 22),
    (3, "regional_terrain", 23),
    (4, "regional_land_improvement", 25),
    (5, "regional_school_proximity", 27),
    (5, "regional_market_proximity", 28),
    (5, "regional_park_proximity", 29),
    (5, "regional_tourism_proximity", 30),
    (5, "regional_parking_convenience", 31),
    (5, "regional_service_facility_proximity", 32),
    (6, "regional_utility_facility_proximity", 34),
    (6, "regional_funeral_facility_proximity", 35),
    (6, "regional_waste_facility_proximity", 36),
    (7, "regional_pollution_proximity", 38),
    (8, "regional_other", 40),
]
_TABLE51_SUBTOTAL_ROWS = {1: 11, 2: 18, 3: 24, 4: 26, 5: 33, 6: 37, 7: 39, 8: 41}
_TABLE51_GRAND_TOTAL_ROW = 42


def build_table51_mapping() -> list:
    sheet = "表5-1區域因素明細表(住)"
    xs, ys = _sheet_grid(
        "表5影響地價區域因素分析明細表(住宅用地).xlsx", sheet, "shulin_table51_blank_v1.pdf",
        num_cols=13, num_rows=45,
    )

    entries = [
        _entry("segment_code[base]", sheet, "C3:D3", xs, ys, align="center"),
        _entry("segment_code[comp1]", sheet, "E3:G3", xs, ys, align="center"),
        _entry("segment_code[comp2]", sheet, "H3:J3", xs, ys, align="center"),
        _entry("segment_code[comp3]", sheet, "K3:M3", xs, ys, align="center"),
    ]
    for _cat, field_id, row in _TABLE51_FACTOR_ROWS:
        entries.append(_entry(f"{field_id}.base_grade", sheet, f"C{row}:C{row}", xs, ys, align="center"))
        entries.append(_entry(f"{field_id}.comp1_grade", sheet, f"E{row}:E{row}", xs, ys, align="center"))
        entries.append(_entry(f"{field_id}.comp1_pct", sheet, f"G{row}:G{row}", xs, ys, align="right", formatter="pct_signed"))
        entries.append(_entry(f"{field_id}.comp2_grade", sheet, f"H{row}:H{row}", xs, ys, align="center"))
        entries.append(_entry(f"{field_id}.comp2_pct", sheet, f"J{row}:J{row}", xs, ys, align="right", formatter="pct_signed"))
        entries.append(_entry(f"{field_id}.comp3_grade", sheet, f"K{row}:K{row}", xs, ys, align="center"))
        entries.append(_entry(f"{field_id}.comp3_pct", sheet, f"M{row}:M{row}", xs, ys, align="right", formatter="pct_signed"))

    for cat, row in _TABLE51_SUBTOTAL_ROWS.items():
        entries.append(_entry(f"category{cat}.comp1_subtotal_pct", sheet, f"E{row}:F{row}", xs, ys, align="right", formatter="pct_signed_no_unit"))
        entries.append(_entry(f"category{cat}.comp2_subtotal_pct", sheet, f"H{row}:I{row}", xs, ys, align="right", formatter="pct_signed_no_unit"))
        entries.append(_entry(f"category{cat}.comp3_subtotal_pct", sheet, f"K{row}:L{row}", xs, ys, align="right", formatter="pct_signed_no_unit"))

    row = _TABLE51_GRAND_TOTAL_ROW
    entries.append(_entry("grand_total.comp1_pct", sheet, f"E{row}:F{row}", xs, ys, align="right", formatter="pct_signed_no_unit"))
    entries.append(_entry("grand_total.comp2_pct", sheet, f"H{row}:I{row}", xs, ys, align="right", formatter="pct_signed_no_unit"))
    entries.append(_entry("grand_total.comp3_pct", sheet, f"K{row}:L{row}", xs, ys, align="right", formatter="pct_signed_no_unit"))
    return entries


# ---------------------------------------------------------------------------
# TABLE4 (表4比較法調查估價表) -- 0基本資料 (rows 5-8) + 20 individual
# factor slots (rows 9-28, incl. FAR at row 26 and catch-all "其他" at row
# 28) + summary block (rows 29-32).
# ---------------------------------------------------------------------------

# field_id (engine/individual_factor_catalog.py's own catalog) -> Table4 row.
# Verified cell-by-cell against the real 表4比較法調查估價表 sheet (docs/
# phase7/official_six_page_pdf_e1.md Task 1): rows 9-25 and 27 are the 18
# non-FAR/non-catch-all IND-* rule-derived factors (in official form
# numbering order 7-23,25), row 26 (numbered item "24容積率(%)") is the FAR
# special-policy slot (individual_floor_area_ratio -- deliberately no rule
# record, see engine/individual_factor_catalog.py's FAR_FIELD_ID), row 28
# ("6其他", unnumbered catch-all) is individual_other.
_TABLE4_INDIVIDUAL_ROWS = {
    "individual_land_area": 9,
    "individual_land_width": 10,
    "individual_land_depth": 11,
    "individual_land_shape": 12,
    "individual_street_frontage": 13,
    "individual_terrain_form4": 14,
    "individual_road_type": 15,
    "individual_frontage_road_width": 16,
    "individual_school_proximity": 17,
    "individual_market_proximity": 18,
    "individual_park_proximity": 19,
    "individual_station_proximity": 20,
    "individual_commercial_district_proximity": 21,
    "individual_nuisance_facility": 22,
    "individual_parking_convenience": 23,
    "individual_zoning_designation": 24,
    "individual_building_coverage_ratio": 25,
    "individual_floor_area_ratio": 26,  # FAR special policy -- see Task 6/7
    "individual_construction_restriction": 27,
    "individual_other": 28,
}

# Per-comparable column groups (1-indexed openpyxl columns): condition
# (base's own raw value / comparable's own raw value) spans 3 cols,
# diff%(P00N only) is the 1 column right after.
_TABLE4_COLS = {
    "base": {"cond": ("D", "F")},
    "comp1": {"cond": ("G", "I"), "pct": "J"},
    "comp2": {"cond": ("K", "M"), "pct": "N"},
    "comp3": {"cond": ("O", "Q"), "pct": "R"},
}


def build_table4_mapping() -> list:
    sheet = "表4比較法調查估價表"
    xs, ys = _sheet_grid(
        "表4比較法調查估價表.xlsx", sheet, "shulin_table4_blank_v1.pdf", num_cols=18, num_rows=36,
    )

    entries = []
    # Header / case-fixed fields.
    # NOTE (fixed after visual smoke-test verification): K1="估價基準日："
    # and O1="案號：" are the LABEL cells themselves (single, unmerged) --
    # appraisal_base_date's real value slot is the merge right after K1
    # (L1:N1); case_no's is the (unmerged, but blank/available) columns
    # right after O1 (P1:R1). The original K1:N1/O1:R1 mapping wrongly
    # started each rect AT its own label cell, visually swapping which
    # value seemed to belong to which label.
    entries.append(_entry("appraisal_base_date", sheet, "L1:N1", xs, ys, align="center"))
    entries.append(_entry("case_no", sheet, "P1:R1", xs, ys, align="center"))
    entries.append(_entry("base_parcel_serial_number", sheet, "H2:I2", xs, ys, align="center"))
    entries.append(_entry("comp1_instance_no", sheet, "L2:M2", xs, ys, align="center"))
    entries.append(_entry("comp2_instance_no", sheet, "P2:Q2", xs, ys, align="center"))
    # NOTE: comp3 has no 4th "實例編號" header slot on this 3-comparable-wide
    # official form beyond O2 (which is the "比較標的3" label itself, not a
    # separate 編號 box) -- left unmapped rather than guessed.

    # 0基本資料 rows: 5=土地正常單價, 6=交易日期+調整百分率, 7=調整至估價基準日單價, 8=地價區段號+區域因素調整百分率
    for cid, (col_lo, col_hi) in (("comp1", ("G", "J")), ("comp2", ("K", "N")), ("comp3", ("O", "R"))):
        entries.append(_entry(f"{cid}.land_normal_price_raw", sheet, f"{col_lo}5:{col_hi}5", xs, ys, align="right", formatter="int_no_unit"))
    entries.append(_entry("comp1.transaction_date_raw", sheet, "G6:I6", xs, ys, align="center"))
    entries.append(_entry("comp1.price_date_adjustment_pct_raw", sheet, "J6:J6", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("comp2.transaction_date_raw", sheet, "K6:M6", xs, ys, align="center"))
    entries.append(_entry("comp2.price_date_adjustment_pct_raw", sheet, "N6:N6", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("comp3.transaction_date_raw", sheet, "O6:Q6", xs, ys, align="center"))
    entries.append(_entry("comp3.price_date_adjustment_pct_raw", sheet, "R6:R6", xs, ys, align="right", formatter="pct_no_unit"))
    for cid, (col_lo, col_hi) in (("comp1", ("G", "J")), ("comp2", ("K", "N")), ("comp3", ("O", "R"))):
        entries.append(_entry(f"{cid}.adjusted_price_raw", sheet, f"{col_lo}7:{col_hi}7", xs, ys, align="right", formatter="int_no_unit"))
    entries.append(_entry("base.segment_code", sheet, "D8:F8", xs, ys, align="center"))
    entries.append(_entry("comp1.segment_code", sheet, "G8:I8", xs, ys, align="center"))
    entries.append(_entry("comp1.regional_adjustment_pct", sheet, "J8:J8", xs, ys, align="right", formatter="pct_signed"))
    entries.append(_entry("comp2.segment_code", sheet, "K8:M8", xs, ys, align="center"))
    entries.append(_entry("comp2.regional_adjustment_pct", sheet, "N8:N8", xs, ys, align="right", formatter="pct_signed"))
    entries.append(_entry("comp3.segment_code", sheet, "O8:Q8", xs, ys, align="center"))
    entries.append(_entry("comp3.regional_adjustment_pct", sheet, "R8:R8", xs, ys, align="right", formatter="pct_signed"))

    # 20 individual factor rows.
    for field_id, row in _TABLE4_INDIVIDUAL_ROWS.items():
        entries.append(_entry(f"individual.{field_id}.base_raw", sheet, f"D{row}:F{row}", xs, ys, font_size=7.5))
        entries.append(_entry(f"individual.{field_id}.comp1_raw", sheet, f"G{row}:I{row}", xs, ys, font_size=7.5))
        entries.append(_entry(f"individual.{field_id}.comp1_pct", sheet, f"J{row}:J{row}", xs, ys, align="right", formatter="pct_signed"))
        entries.append(_entry(f"individual.{field_id}.comp2_raw", sheet, f"K{row}:M{row}", xs, ys, font_size=7.5))
        entries.append(_entry(f"individual.{field_id}.comp2_pct", sheet, f"N{row}:N{row}", xs, ys, align="right", formatter="pct_signed"))
        entries.append(_entry(f"individual.{field_id}.comp3_raw", sheet, f"O{row}:Q{row}", xs, ys, font_size=7.5))
        entries.append(_entry(f"individual.{field_id}.comp3_pct", sheet, f"R{row}:R{row}", xs, ys, align="right", formatter="pct_signed"))

    # Summary block: row29=合計(individual total), row30 left-half=調整絕對值加總,
    # row31 left-half=試算價格/right-half=權重, row32=比準地比較價格 (single wide cell).
    entries.append(_entry("comp1.individual_adjustment_total_pct", sheet, "G29:J29", xs, ys, align="right", formatter="pct_signed"))
    entries.append(_entry("comp2.individual_adjustment_total_pct", sheet, "K29:N29", xs, ys, align="right", formatter="pct_signed"))
    entries.append(_entry("comp3.individual_adjustment_total_pct", sheet, "O29:R29", xs, ys, align="right", formatter="pct_signed"))
    entries.append(_entry("comp1.adjustment_abs_sum", sheet, "G30:H30", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("comp2.adjustment_abs_sum", sheet, "K30:L30", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("comp3.adjustment_abs_sum", sheet, "O30:P30", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("comp1.trial_price", sheet, "G31:H31", xs, ys, align="right", formatter="int_no_unit"))
    entries.append(_entry("comp2.trial_price", sheet, "K31:L31", xs, ys, align="right", formatter="int_no_unit"))
    entries.append(_entry("comp3.trial_price", sheet, "O31:P31", xs, ys, align="right", formatter="int_no_unit"))
    entries.append(_entry("comp1.weight_pct", sheet, "I31:J31", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("comp2.weight_pct", sheet, "M31:N31", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("comp3.weight_pct", sheet, "Q31:R31", xs, ys, align="right", formatter="pct_no_unit"))
    entries.append(_entry("base_comparison_price", sheet, "G32:R32", xs, ys, align="right", formatter="int_no_unit"))
    return entries


def build_all() -> None:
    mappings = {
        "shulin_table3_mapping.json": build_table3_mapping(),
        "shulin_table51_mapping.json": build_table51_mapping(),
        "shulin_table4_mapping.json": build_table4_mapping(),
    }
    for filename, entries in mappings.items():
        out_path = os.path.join(_TEMPLATE_DIR, filename)
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(entries, f, ensure_ascii=False, indent=1)
        print(f"{filename}: {len(entries)} fields")


if __name__ == "__main__":
    build_all()
