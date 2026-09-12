# -*- coding: utf-8 -*-
"""
SUPPLEMENTAL-JSON-EXCEL-EXPORT-H1 tests.

Minimal, high-value validation only (per this round's own explicit
instruction) -- reuses the same moto+real-handler pattern established in
C1/D1/E1/F1's own test files. export/*.py never imports pdf_handler.py or
weasyprint (only pdf/shulin_official_pdf_renderer.py, which is PyMuPDF-
only), so none of the sys.modules-stubbing machinery those other test
files needed applies here.
"""
from __future__ import annotations

import datetime
import hashlib
import json
import os
import sys

import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO_ROOT, "backend", "handlers"))
sys.path.insert(0, os.path.join(REPO_ROOT, "engine"))
sys.path.insert(0, os.path.join(REPO_ROOT, "providers"))
sys.path.insert(0, os.path.join(REPO_ROOT, "pdf"))
sys.path.insert(0, os.path.join(REPO_ROOT, "data", "competition_cases", "shulin_residential_2026"))
sys.path.insert(0, REPO_ROOT)

try:
    from moto import mock_aws
    import boto3
except ImportError:
    pytest.skip("moto/boto3 not available", allow_module_level=True)

try:
    import openpyxl
except ImportError:
    pytest.skip("openpyxl not available", allow_module_level=True)

try:
    import fitz
except ImportError:
    pytest.skip("PyMuPDF not available", allow_module_level=True)

import segment_table3_fixtures as fx  # noqa: E402
import segment_table4_fixtures as fx4  # noqa: E402
from domain.models import CaseRulePackage  # noqa: E402

PROFILE_ID = "shulin_residential_2026"
SHULIN_RULES_DIR = os.path.join(REPO_ROOT, "data", "rules", "competition", "shulin_residential_2026")


def _load_rules(name):
    with open(os.path.join(SHULIN_RULES_DIR, name), encoding="utf-8") as f:
        return json.load(f)["rules"]


REGIONAL_RULES = _load_rules("regional_rules.json")
INDIVIDUAL_RULES = _load_rules("individual_rules.json")

TABLE3_SOURCE_XLSX = os.path.join(REPO_ROOT, "data", "sources", "competition", "shulin_residential_2026", "表3地價區段勘查表.xlsx")
TABLE51_SOURCE_XLSX = os.path.join(REPO_ROOT, "data", "sources", "competition", "shulin_residential_2026", "表5影響地價區域因素分析明細表(住宅用地).xlsx")
TABLE4_SOURCE_XLSX = os.path.join(REPO_ROOT, "data", "sources", "competition", "shulin_residential_2026", "表4比較法調查估價表.xlsx")


@pytest.fixture(autouse=True)
def _restore_data_provider_mode_for_later_modules(monkeypatch):
    yield
    import collect_data
    import importlib
    monkeypatch.setenv("DATA_PROVIDER_MODE", "mock")
    importlib.reload(collect_data)


def _ensure_mock_collect_data(monkeypatch):
    import collect_data
    import importlib
    monkeypatch.setenv("DATA_PROVIDER_MODE", "mock")
    importlib.reload(collect_data)
    return collect_data


@pytest.fixture()
def ddb_env(monkeypatch, tmp_path):
    monkeypatch.setenv("CASES_TABLE_NAME", "test-table-h1")
    monkeypatch.setenv("AWS_DEFAULT_REGION", "ap-northeast-1")
    monkeypatch.setenv("DATASET_REGISTRY_DB_PATH", str(tmp_path / "dataset_registry.sqlite3"))
    with mock_aws():
        ddb = boto3.client("dynamodb", region_name="ap-northeast-1")
        ddb.create_table(
            TableName="test-table-h1",
            KeySchema=[{"AttributeName": "PK", "KeyType": "HASH"}, {"AttributeName": "SK", "KeyType": "RANGE"}],
            AttributeDefinitions=[{"AttributeName": "PK", "AttributeType": "S"},
                                   {"AttributeName": "SK", "AttributeType": "S"}],
            BillingMode="PAY_PER_REQUEST",
        )
        from _aws_mock_reset import reset_cached_aws_module_state
        reset_cached_aws_module_state()
        yield


def _make_shulin_case(cases_module, case_no):
    body = {
        "case_no": case_no, "segment_code": "P001-00", "city": "新北市", "district": fx.DISTRICT,
        "land_use_type": fx.LAND_USE_TYPE, "appraisal_period": fx.APPRAISAL_PERIOD,
        "appraisal_base_date": fx.APPRAISAL_PERIOD, "segment_scope": "P001-00 比準地區段",
        "base_parcel_id": fx.P001_SEGMENT_META["parcel_ids"][0],
        "comparable_ids": ["P002-00", "P003-00", "P004-00"],
        "rule_profile_id": PROFILE_ID, "segments": fx.segment_map_body(),
    }
    return cases_module.create_case({"body": json.dumps(body)}, None)


def _collect_segment(collect_data_module, case_no, segment_code):
    cd_body = {"competition_provided_factors": fx.SEGMENT_TABLE3_FACTORS[segment_code]}
    if segment_code in fx4.SEGMENT_TABLE4_TRANSACTION:
        cd_body["competition_provided_transaction"] = fx4.SEGMENT_TABLE4_TRANSACTION[segment_code]
    return collect_data_module.collect_data(
        {"pathParameters": {"id": case_no, "segment_code": segment_code}, "body": json.dumps(cd_body)}, None,
    )


def _confirm_shulin_package(case_no, package_id="PKG-H1"):
    from case_rule_repository import DynamoCaseRuleRepository
    repo = DynamoCaseRuleRepository()
    pkg = CaseRulePackage(
        case_id=case_no, package_id=package_id, rule_version="h1-test-v1",
        source_document="評價基準明細表.pdf", source_type="MANUAL_CSV_INGEST",
        regional_rules=REGIONAL_RULES, individual_rules=INDIVIDUAL_RULES,
        metadata={"competition_profile": {
            "profile_id": PROFILE_ID, "district": "樹林區", "land_use_type": "普通住宅用地",
            "source_document": "評價基準明細表.pdf",
            "source_sha256": "a7574aaf56b546737df8b3b34459e764523be77ae4af25490d617a41a33ed09c",
        }},
        created_at=datetime.datetime.now(datetime.timezone.utc),
    )
    repo.save_candidate(pkg)
    repo.confirm(case_no, package_id, confirmed_by="tester")


def _full_shulin_setup(monkeypatch, case_no):
    import cases
    collect_data = _ensure_mock_collect_data(monkeypatch)
    create_resp = _make_shulin_case(cases, case_no)
    assert create_resp["statusCode"] in (200, 201), create_resp["body"]
    for code in ("P001-00", "P002-00", "P003-00", "P004-00"):
        r = _collect_segment(collect_data, case_no, code)
        assert r["statusCode"] == 200, r["body"]
    _confirm_shulin_package(case_no)


def _sha256(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


@pytest.fixture()
def bundle(ddb_env, monkeypatch):
    case_no = "H1-EXPORT-001"
    _full_shulin_setup(monkeypatch, case_no)
    from export.bundle_builder import build_case_export_bundle
    return build_case_export_bundle(case_no)


# ---------------------------------------------------------------------------
# 1/2/3. JSON export: produces valid JSON, contains all 4 segments, and
# Table51/Table4 lineage is internally consistent.
# ---------------------------------------------------------------------------

class TestJsonExport:
    def test_1_2_3_json_export_valid_all_segments_and_lineage(self, bundle):
        from export.json_exporter import export_bundle_to_json_bytes
        raw = export_bundle_to_json_bytes(bundle)
        parsed = json.loads(raw)  # must not raise

        assert set(parsed["segments"].keys()) == {"P001-00", "P002-00", "P003-00", "P004-00"}
        assert set(parsed["table3"].keys()) == {"P001-00", "P002-00", "P003-00", "P004-00"}

        assert parsed["table5_1"]["base_segment_code"] == "P001-00"
        t51_by_idx = {c["comparison_index"]: c["comparable_segment_code"] for c in parsed["table5_1"]["comparisons"]}
        assert t51_by_idx == {1: "P002-00", 2: "P003-00", 3: "P004-00"}
        t4_by_idx = {c["comparison_index"]: c["comparable_segment_code"] for c in parsed["table4"]["comparisons"]}
        assert t4_by_idx == t51_by_idx  # same lineage on both sides -- no cross-contamination

        assert parsed["review"] is None
        assert parsed["review_status"] == "NOT_AVAILABLE_FOR_SHULIN_YET"
        assert isinstance(parsed["manual_review_items"], list) and len(parsed["manual_review_items"]) > 0
        assert "provenance" in parsed and isinstance(parsed["provenance"], dict)
        assert parsed["schema_version"] == "1.0"


# ---------------------------------------------------------------------------
# 4/5. All 6 Excel files open cleanly via openpyxl; source XLSX untouched.
# ---------------------------------------------------------------------------

class TestExcelExportIntegrity:
    def test_4_5_all_excel_files_open_and_source_untouched(self, bundle, tmp_path):
        before = {p: _sha256(p) for p in (TABLE3_SOURCE_XLSX, TABLE51_SOURCE_XLSX, TABLE4_SOURCE_XLSX)}

        from export.excel_exporter import export_all_excel_files
        paths = export_all_excel_files(bundle, str(tmp_path))
        assert len(paths) == 6
        for logical_name, path in paths.items():
            wb = openpyxl.load_workbook(path)
            assert wb.sheetnames  # opens without raising
            wb.close()

        after = {p: _sha256(p) for p in before}
        assert before == after


# ---------------------------------------------------------------------------
# 6/7/8. Missing != zero; FAR blank; unconfirmed weight blank -- checked on
# the actual written Table4 Excel cells.
# ---------------------------------------------------------------------------

class TestExcelSafety:
    def test_6_7_8_missing_far_and_weight_stay_blank_in_excel(self, bundle, tmp_path):
        from export.excel_exporter import export_table4_excel
        import json as _json
        mapping_path = os.path.join(REPO_ROOT, "data", "templates", "shulin", "shulin_table4_mapping.json")
        mapping = {e["field_id"]: e for e in _json.load(open(mapping_path, encoding="utf-8"))}

        out_path = str(tmp_path / "table4.xlsx")
        export_table4_excel(bundle, out_path)
        wb = openpyxl.load_workbook(out_path)
        ws = wb["表4比較法調查估價表"]

        far_entry = mapping["individual.individual_floor_area_ratio.comp1_pct"]
        far_cell = far_entry["source_cell"].split(":")[0]
        assert ws[far_cell].value in (None, "")

        weight_entry = mapping["comp1.weight_pct"]
        weight_cell = weight_entry["source_cell"].split(":")[0]
        assert ws[weight_cell].value in (None, "")
        wb.close()


# ---------------------------------------------------------------------------
# 9. P002/P003/P004 fixed transaction data correct in both JSON and Excel
# (Task 20 consistency check).
# ---------------------------------------------------------------------------

class TestFixedTransactionConsistency:
    def test_9_p002_p003_p004_fixed_data_matches_in_json_and_excel(self, bundle, tmp_path):
        expected = {
            "P002-00": ("110年9月14日", 130167, "5.96", 137925),
            "P003-00": ("111年1月11日", 135275, "4.09", 140808),
            "P004-00": ("110年10月29日", 170909, "5.49", 180292),
        }
        t4 = bundle.table4
        by_seg = {c["comparable_segment_code"]: c for c in t4["comparisons"]}
        for seg, (date, price, pct, adjusted) in expected.items():
            comp = by_seg[seg]
            assert comp["transaction_date_raw"] == date
            assert float(comp["land_normal_price_raw"]) == price
            assert float(comp["price_date_adjustment_pct_raw"]) == float(pct)
            assert float(comp["adjusted_price_raw"]) == adjusted

        from export.excel_exporter import export_table4_excel
        out_path = str(tmp_path / "table4.xlsx")
        export_table4_excel(bundle, out_path)
        wb = openpyxl.load_workbook(out_path)
        ws = wb["表4比較法調查估價表"]
        # comp1 (P002-00) land_normal_price_raw cell, cross-checked directly
        # against the mapping this exporter itself uses.
        import json as _json
        mapping_path = os.path.join(REPO_ROOT, "data", "templates", "shulin", "shulin_table4_mapping.json")
        mapping = {e["field_id"]: e for e in _json.load(open(mapping_path, encoding="utf-8"))}
        cell = mapping["comp1.land_normal_price_raw"]["source_cell"].split(":")[0]
        assert ws[cell].value == 130167
        wb.close()


# ---------------------------------------------------------------------------
# 10. Official 6-page PDF smoke -- no regression from H1 (same renderer,
# fed by the SAME bundle).
# ---------------------------------------------------------------------------

class TestOfficialPdfNoRegression:
    def test_10_official_pdf_still_six_pages_from_bundle(self, bundle):
        from export.pdf_export import render_official_pdf_from_bundle
        pdf_bytes = render_official_pdf_from_bundle(bundle)
        assert pdf_bytes.startswith(b"%PDF")
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        assert len(doc) == 6
        order_markers = ["P002-00", "P003-00", "P004-00", "P001-00", "表5-1", "表4"]
        for i, marker in enumerate(order_markers):
            assert marker in doc[i].get_text()
        doc.close()


# ---------------------------------------------------------------------------
# Bundle ZIP: sanity that the zip contains json + 6 excel + pdf.
# ---------------------------------------------------------------------------

class TestZipBundle:
    def test_11_zip_bundle_contains_json_excel_and_pdf(self, bundle):
        import zipfile
        import io
        from export.zip_bundle import build_export_zip_bytes, bundle_zip_filename
        zip_bytes = build_export_zip_bytes(bundle, include_pdf=True)
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
        names = zf.namelist()
        assert any(n.endswith("_data.json") for n in names)
        assert sum(1 for n in names if n.startswith("excel/") and n.endswith(".xlsx")) == 6
        assert "official_6_page.pdf" in names
        assert bundle_zip_filename(bundle.case_no).startswith("case_")
